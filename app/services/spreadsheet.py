from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook


ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}
MAX_RECORDS = 10_000

# Nomes das colunas que a planilha precisa ter na primeira linha.
REQUIRED_COLUMNS = {
    "prova": "prova",
    "nomeue": "nome_ue",
    "nomealuno": "nome_aluno",
    "turma": "turma",
    "alunora": "ra",
    "componente": "componente",
    "proficiencia": "proficiencia",
    "nivel": "nivel",
}


class SpreadsheetReadError(Exception):
    pass


def normalize_column_name(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "").strip())
    without_accents = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return "".join(char.lower() for char in without_accents if char.isalnum())


def normalize_ra(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\.0+$", "", str(value or "").strip())


def normalize_score(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace(",", ".")), 2)
    except ValueError:
        return None


def normalize_semester(value: object) -> int:
    match = re.search(r"[1-2]", str(value or ""))
    if not match:
        raise SpreadsheetReadError(
            "A coluna SEMESTRE deve informar o 1º ou o 2º semestre."
        )
    return int(match.group(0))


def read_spreadsheet(contents: bytes) -> list[dict]:
    try:
        workbook = load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=True,
            keep_vba=False,
        )
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows)
    except Exception as exc:
        raise SpreadsheetReadError("Não foi possível ler a planilha.") from exc

    # Guarda a posicao de cada coluna para facilitar a leitura das linhas.
    columns_by_key = {
        normalize_column_name(header): index for index, header in enumerate(headers)
    }
    missing = [key for key in REQUIRED_COLUMNS if key not in columns_by_key]
    if "semestre" not in columns_by_key:
        missing.append("semestre")
    if missing:
        raise SpreadsheetReadError(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}."
        )

    records: list[dict] = []
    for row in rows:
        def cell_value(key: str) -> object:
            index = columns_by_key[key]
            return row[index] if index < len(row) else None

        ra = normalize_ra(cell_value("alunora"))
        if not ra:
            continue
        if len(records) >= MAX_RECORDS:
            raise SpreadsheetReadError("A planilha ultrapassa 10.000 registros.")

        record = {
            output_name: str(cell_value(key) or "").strip()
            for key, output_name in REQUIRED_COLUMNS.items()
        }
        # O banco usa o nome bimestre por compatibilidade, mas a planilha informa semestre.
        record["bimestre"] = normalize_semester(cell_value("semestre"))
        record["ra"] = ra
        record["proficiencia"] = normalize_score(
            cell_value("proficiencia")
        )
        records.append(record)

    workbook.close()
    if not records:
        raise SpreadsheetReadError("A planilha não possui registros válidos.")
    return records
